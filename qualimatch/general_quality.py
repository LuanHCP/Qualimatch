from __future__ import annotations

"""
qualimatch/general_quality.py
------------------------------
Módulo de Qualidade Geral: leitura das planilhas de Traços (inventário de
traços aprovados) e de Não Conformidades (controle de RNCs), e toda a lógica
de resumo/filtro usada pelas telas correspondentes.

Este código vivia originalmente dentro de app.py (bloco "Qualidade Geral",
reintroduzido ali em 28/08/2026 como correção emergencial). Foi movido para
cá em 31/08/2026 para seguir o mesmo padrão do resto do projeto — onde
app.py contém apenas rotas finas que delegam para qualimatch/*.py — sem
nenhuma mudança de comportamento: é uma extração literal, função por
função, validada linha a linha contra os mesmos dados reais usados para
corrigir os bugs de retenção/leitura ao vivo (ver histórico do projeto).

Convenção mantida de propósito: os nomes com "_" no início (ex.: _qg_norm)
já eram assim quando viviam em app.py. Não foram renomeados aqui para
manter esta extração o mais literal possível; um "des-underscoring" das
funções realmente públicas deste módulo (as usadas pelas rotas em app.py)
é uma limpeza válida para uma próxima etapa, mas fica fora do escopo desta
extração.
"""

import os
import re
import shutil
import time
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from urllib.parse import unquote, urljoin, urlparse
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

from .config import DATA_DIR
from .database import row, transaction

_QG_DIR = DATA_DIR / "quality_general"
_QG_TRACE_DEFAULT = _QG_DIR / "inventario_tracos.xlsx"
_QG_NC_DEFAULT = _QG_DIR / "controle_nao_conformidades.xlsx"
_QG_TRACE_SETTING = "general_quality_inventory_path"
_QG_NC_SETTING = "general_quality_nc_path"
_QG_TRACE_SHEETS = ("CBUQ", "CONCRETO", "MRAF", "TSD", "PMF", "BGS", "BGTC", "MACADAME")
_qg_trace_cache = {"path": "", "mtime_ns": None, "size": None, "records": []}
_qg_nc_cache = {"path": "", "mtime_ns": None, "size": None, "records": []}
_qg_source_copy_cache: dict[str, tuple[str, int, int]] = {}


def _qg_norm(value) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^A-Za-z0-9]+", " ", text).strip().upper()
    return re.sub(r"\s+", " ", text)


def _qg_clean(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value)).strip()


def _qg_get_setting(key: str) -> str:
    try:
        item = row("SELECT value FROM settings WHERE key=?", (key,))
        return str(item.get("value") or "").strip() if item else ""
    except Exception:
        return ""


def _qg_set_setting(key: str, value: str) -> None:
    with transaction() as con:
        con.execute(
            "INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
            (key, value),
        )


def _qg_path(setting_key: str, default: Path) -> Path:
    configured = _qg_get_setting(setting_key)
    return Path(configured).expanduser() if configured else default


def _qg_normalize_source_path(value: str) -> Path:
    text = str(value or "").strip().strip('"').strip("'").strip()
    if not text:
        return Path()
    parsed = urlparse(text)
    if parsed.scheme.lower() in {"http", "https"}:
        raise ValueError(
            "Cole o caminho do arquivo sincronizado no Windows (C:\\...), "
            "não o link https:// do SharePoint."
        )
    if parsed.scheme.lower() == "file":
        local = unquote(parsed.path)
        if parsed.netloc:
            local = f"//{parsed.netloc}{local}"
        elif re.match(r"^/[A-Za-z]:", local):
            local = local[1:]
        text = local
    return Path(os.path.expandvars(text)).expanduser()


def _qg_validate_source_file(path: Path) -> None:
    if not path.exists() or not path.is_file():
        raise ValueError(
            "Não encontrei a planilha nesse caminho. Confirme se o arquivo está sincronizado "
            "e disponível localmente no OneDrive/SharePoint."
        )
    try:
        with ZipFile(path) as archive:
            archive.getinfo("xl/workbook.xml")
    except (BadZipFile, KeyError) as exc:
        raise ValueError("O arquivo informado não é uma planilha Excel válida.") from exc
    except (OSError, PermissionError) as exc:
        raise ValueError(
            "Não consegui ler a planilha sincronizada. Aguarde o OneDrive concluir o download "
            "ou marque o arquivo como 'Sempre manter neste dispositivo'."
        ) from exc


def _qg_set_path(setting_key: str, value: str, default: Path) -> Path:
    text = str(value or "").strip().strip('"').strip("'").strip()
    if not text:
        _qg_set_setting(setting_key, "")
        return default
    p = _qg_normalize_source_path(text)
    if p.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("A fonte precisa ser uma planilha .xlsx ou .xlsm.")
    _qg_validate_source_file(p)
    _qg_set_setting(setting_key, str(p))
    return p


def _qg_source_info(path: Path, default: Path) -> dict:
    exists = path.exists() and path.is_file()
    st = path.stat() if exists else None
    try:
        managed = path.resolve() == default.resolve()
    except Exception:
        managed = False
    return {
        "path": str(path),
        "exists": exists,
        "file_name": path.name,
        "size": st.st_size if st else 0,
        "modified_at": datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds") if st else None,
        "mode": "managed_upload" if managed else "external_path",
    }


def _qg_excel_date(value):
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            from openpyxl.utils.datetime import from_excel
            d = from_excel(value)
            return (d.date() if isinstance(d, datetime) else d).isoformat()
        except Exception:
            try:
                return (date(1899, 12, 30) + timedelta(days=int(value))).isoformat()
            except Exception:
                return None
    text = _qg_clean(value)
    if not text:
        return None
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", text)
    if m:
        try:
            return date(int(m.group(1)), int(m.group(2)), int(m.group(3))).isoformat()
        except ValueError:
            return text
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text


def _qg_money(value) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    text = re.sub(r"[^0-9.\-]", "", text)
    try:
        return float(text or 0)
    except ValueError:
        return 0.0


def _qg_headers(values) -> dict[str, int]:
    return {_qg_norm(v): i for i, v in enumerate(values) if _qg_norm(v)}


def _qg_pick(values, headers: dict[str, int], *names):
    for name in names:
        i = headers.get(_qg_norm(name))
        if i is not None and i < len(values):
            return values[i]
    return None


def _qg_workbook_link_base(path: Path) -> str:
    """Lê a pasta-base que o Excel usa para hyperlinks relativos do SharePoint."""
    try:
        with ZipFile(path) as archive:
            workbook_xml = archive.read("xl/workbook.xml").decode("utf-8", "replace")
    except (BadZipFile, KeyError, OSError):
        return ""
    match = re.search(r"<[^>]*absPath\b[^>]*\burl=([\"'])(.*?)\1", workbook_xml, re.I | re.S)
    if not match:
        return ""
    base = match.group(2).replace("&amp;", "&").strip()
    return base if base.lower().startswith(("http://", "https://")) else ""


def _qg_resolve_link(target: str, link_base: str = "") -> str:
    value = str(target or "").strip()
    if not value or value.lower().startswith(("http://", "https://")):
        return value
    if link_base and not urlparse(value).scheme:
        resolved = urljoin(link_base.rstrip("/") + "/", value)
        if resolved.lower().startswith(("http://", "https://")):
            return resolved
    return value


def _qg_link(ws, row_num: int, headers: dict[str, int], *names, link_base: str = "") -> str:
    for name in names:
        i = headers.get(_qg_norm(name))
        if i is None:
            continue
        cell = ws.cell(row=row_num, column=i + 1)
        if cell.hyperlink and cell.hyperlink.target:
            return _qg_resolve_link(cell.hyperlink.target, link_base)
        val = _qg_clean(cell.value)
        if val.startswith(("http://", "https://")):
            return val
    return ""


def _qg_stable_source_copy(source: Path, default: Path, cache_name: str) -> Path:
    """Materializa uma cópia estável de fontes sincronizadas antes de abrir o XLSX."""
    try:
        if source.resolve() == default.resolve():
            return source
    except OSError:
        pass
    source_key = str(source.resolve())
    source_stat = source.stat()
    signature = (source_key, source_stat.st_mtime_ns, source_stat.st_size)
    suffix = source.suffix.lower() if source.suffix.lower() in {".xlsx", ".xlsm"} else ".xlsx"
    cache_dir = _QG_DIR / "cache"
    cached = cache_dir / f"{cache_name}{suffix}"
    if _qg_source_copy_cache.get(cache_name) == signature and cached.exists():
        return cached
    cache_dir.mkdir(parents=True, exist_ok=True)
    tmp = cached.with_name(f".{cached.name}.{os.getpid()}.tmp")
    last_error = None
    for _ in range(5):
        try:
            before = source.stat()
            with source.open("rb") as src, tmp.open("wb") as dst:
                shutil.copyfileobj(src, dst, length=1024 * 1024)
            after = source.stat()
            if (before.st_mtime_ns, before.st_size) != (after.st_mtime_ns, after.st_size):
                time.sleep(.25)
                continue
            os.replace(tmp, cached)
            _qg_source_copy_cache[cache_name] = (source_key, after.st_mtime_ns, after.st_size)
            return cached
        except (OSError, PermissionError) as exc:
            last_error = exc
            time.sleep(.25)
        finally:
            tmp.unlink(missing_ok=True)
    raise ValueError(
        "Não consegui criar uma cópia estável da planilha sincronizada. Aguarde o OneDrive/SharePoint "
        "terminar a sincronização e tente novamente."
    ) from last_error


def _qg_trace_group(status: str) -> str:
    s = _qg_norm(status)
    if s == "VENCIDO":
        return "EXPIRED"
    if s == "APROVADO":
        return "APPROVED"
    if s in {"RESTRITO", "APROV COM RESSALVAS", "APROVADO COM RESSALVAS"}:
        return "RESTRICTED"
    if s in {"REPROVADO", "NAO CONFORME", "NAO CONFORMIDADE"}:
        return "REJECTED"
    if s in {"EM ANALISE", "SOLICITADO", "AGUARDANDO"}:
        return "ANALYSIS"
    if "CANCEL" in s:
        return "CANCELLED"
    return "OTHER"


def _qg_refresh_expiration(item: dict) -> dict:
    x = dict(item)
    source_status = str(x.get("source_status") or x.get("status") or "").strip()
    x["source_status"] = source_status
    expiration = str(x.get("expiration_date") or "").strip()
    days = None
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", expiration):
        try:
            days = (date.fromisoformat(expiration) - date.today()).days
        except ValueError:
            pass
    x["days_to_expire"] = days
    x["expired"] = days is not None and days < 0
    if x["expired"]:
        x["status"] = "VENCIDO"
        x["status_group"] = "EXPIRED"
    else:
        x["status"] = source_status
        x["status_group"] = _qg_trace_group(source_status)
    return x


def _qg_find_header(ws, required_sets, max_rows=15):
    for ridx in range(1, min(ws.max_row, max_rows) + 1):
        values = [ws.cell(ridx, c).value for c in range(1, ws.max_column + 1)]
        headers = _qg_headers(values)
        if all(any(_qg_norm(name) in headers for name in group) for group in required_sets):
            return ridx, headers
    return None, {}


def _qg_load_traces(force=True) -> list[dict]:
    p = _qg_path(_QG_TRACE_SETTING, _QG_TRACE_DEFAULT)
    if not p.exists() or not p.is_file():
        return []
    st = p.stat()
    key = str(p.resolve())
    if (
        not force
        and _qg_trace_cache["path"] == key
        and _qg_trace_cache["mtime_ns"] == st.st_mtime_ns
        and _qg_trace_cache["size"] == st.st_size
    ):
        return [_qg_refresh_expiration(x) for x in _qg_trace_cache["records"]]
    read_path = _qg_stable_source_copy(p, _QG_TRACE_DEFAULT, "inventario_tracos_live")
    link_base = _qg_workbook_link_base(read_path)
    try:
        wb = load_workbook(read_path, read_only=False, data_only=True)
    except Exception as exc:
        raise ValueError("A planilha configurada da Qualidade Geral não é um arquivo Excel válido.") from exc
    records = []
    try:
        for sheet in _QG_TRACE_SHEETS:
            actual = next((n for n in wb.sheetnames if _qg_norm(n) == _qg_norm(sheet)), None)
            if not actual:
                continue
            ws = wb[actual]
            header_row, headers = _qg_find_header(ws, [
                ("TRAÇO", "TRACO"),
                ("STATUS",),
            ])
            if not header_row:
                continue
            for ridx in range(header_row + 1, ws.max_row + 1):
                vals = [ws.cell(ridx, c).value for c in range(1, ws.max_column + 1)]
                trace = _qg_clean(_qg_pick(vals, headers, "TRAÇO", "TRACO"))
                status = _qg_clean(_qg_pick(vals, headers, "STATUS"))
                code = _qg_clean(_qg_pick(vals, headers, "CÓD INT LAB CCR", "COD INT LAB CCR", "CÓDIGO INTERNO LABORATÓRIO CCR"))
                if not trace or not status:
                    continue
                approval_date = _qg_excel_date(_qg_pick(vals, headers, "DATA APROVAÇÃO", "DATA DA APROVAÇÃO"))
                expiration_date = _qg_excel_date(_qg_pick(vals, headers, "DATA VENCIMENTO"))
                item = {
                    "sheet": actual,
                    "source_row": ridx,
                    "code": code,
                    "external_code": _qg_clean(_qg_pick(vals, headers, "CÓDIDO EXTERNO - CONTRATADA", "CÓDIGO EXTERNO - CONTRATADA")),
                    "trace": trace,
                    "category": sheet,
                    "trace_type": _qg_clean(_qg_pick(vals, headers, "TIPO DE TRAÇO", "TIPO DE TRACO")) or sheet,
                    "classification": _qg_clean(_qg_pick(vals, headers, "CLASSIFICAÇÃO", "CLASSIFICACAO")),
                    "contractor": _qg_clean(_qg_pick(vals, headers, "CONSTRUTORA", "EMPRESA")),
                    "contract": _qg_clean(_qg_pick(vals, headers, "Nº CONTRATO", "N° CONTRATO", "CONTRATO")),
                    "work": _qg_clean(_qg_pick(vals, headers, "OBRA")),
                    "application": _qg_clean(_qg_pick(vals, headers, "LOCAL DE APLICAÇÃO", "APLICABILIDADE")),
                    "plant": _qg_clean(_qg_pick(vals, headers, "USINA")),
                    "plant_type": _qg_clean(_qg_pick(vals, headers, "TIPO USINA")),
                    "city": _qg_clean(_qg_pick(vals, headers, "CIDADE LOCALIZADA A UNIDADE", "CIDADE LOCALIZADA  A UNIDADE")),
                    "aggregate_origin": _qg_clean(_qg_pick(vals, headers, "ORIGEM DO AGREGADO")),
                    "binder_supplier": _qg_clean(_qg_pick(vals, headers, "FORNECEDOR LIGANTE")),
                    "binder_type": _qg_clean(_qg_pick(vals, headers, "TIPO LIGANTE", "TIPO CAP", "LIGANTE")),
                    "approval_code": _qg_clean(_qg_pick(vals, headers, "CÓDIGO DE APROVAÇÃO GBS", "CODIGO DE APROVACAO GBS")),
                    "source_status": status,
                    "status": status,
                    "status_group": _qg_trace_group(status),
                    "approver": _qg_clean(_qg_pick(vals, headers, "RESPONSÁVEL APROVAÇÃO", "RESPONSAVEL APROVACAO")),
                    "approval_date": approval_date,
                    "expiration_date": expiration_date,
                    "seal": _qg_clean(_qg_pick(vals, headers, "SELO DE APROVAÇÃO", "SELO DE APROVACAO")),
                    "link": _qg_link(
                        ws, ridx, headers, "LINK CARTA TRAÇO", "LINK CARTA TRACO", "LINK",
                        link_base=link_base,
                    ),
                    "observations": _qg_clean(_qg_pick(vals, headers, "OBSERVAÇÕES", "OBSERVAÇÃO", "OBSERVACOES", "OBSERVACAO")),
                }
                records.append(item)
    finally:
        wb.close()
    _qg_trace_cache.update({"path": key, "mtime_ns": st.st_mtime_ns, "size": st.st_size, "records": records})
    return [_qg_refresh_expiration(x) for x in records]


def _qg_trace_filter(items, contractor="ALL", trace_type="ALL", status="ALL", search="", plant="ALL", binder="ALL", approved_only=False):
    cn, tn, sn, pn, bn, q = map(_qg_norm, [contractor, trace_type, status, plant, binder, search])
    out = []
    for x in items:
        if approved_only and x.get("status_group") != "APPROVED":
            continue
        if cn not in {"", "ALL"} and _qg_norm(x.get("contractor")) != cn:
            continue
        if tn not in {"", "ALL"} and _qg_norm(x.get("category")) != tn:
            continue
        if sn not in {"", "ALL"} and _qg_norm(x.get("status")) != sn:
            continue
        if pn not in {"", "ALL"} and _qg_norm(x.get("plant")) != pn:
            continue
        if bn not in {"", "ALL"} and _qg_norm(x.get("binder_type")) != bn:
            continue
        if q:
            hay = _qg_norm(" ".join(str(x.get(k) or "") for k in (
                "code", "external_code", "trace", "contractor", "trace_type", "plant",
                "binder_supplier", "binder_type", "aggregate_origin", "work", "application",
            )))
            if q not in hay:
                continue
        out.append(x)
    return out


def _qg_trace_filters(items):
    def u(k):
        return sorted({str(x.get(k) or "").strip() for x in items if str(x.get(k) or "").strip()}, key=_qg_norm)
    return {
        "contractors": u("contractor"),
        "trace_types": u("category"),
        "statuses": u("status"),
        "plants": u("plant"),
        "binders": u("binder_type"),
    }


def _qg_trace_summary(items):
    groups = Counter(x.get("status_group") or "OTHER" for x in items)
    byc = defaultdict(Counter)
    approved_by = defaultdict(list)
    for x in items:
        c = x.get("contractor") or "Sem contratada"
        byc[c][x.get("status_group") or "OTHER"] += 1
        if x.get("status_group") == "APPROVED":
            approved_by[c].append(x)
    contractor_rows = []
    for c in sorted(byc, key=_qg_norm):
        g = byc[c]
        total = sum(g.values())
        contractor_rows.append({
            "contractor": c, "total": total,
            "approved": g["APPROVED"], "restricted": g["RESTRICTED"],
            "rejected": g["REJECTED"], "analysis": g["ANALYSIS"],
            "expired": g["EXPIRED"], "cancelled": g["CANCELLED"],
            "approved_pct": round(g["APPROVED"] / total * 100, 1) if total else 0,
        })
    type_counts = Counter(x.get("category") or "Outros" for x in items)
    expiring = [x for x in items if x.get("status_group") == "APPROVED" and x.get("days_to_expire") is not None and 0 <= int(x["days_to_expire"]) <= 30]
    return {
        "total": len(items), "approved": groups["APPROVED"], "restricted": groups["RESTRICTED"],
        "rejected": groups["REJECTED"], "analysis": groups["ANALYSIS"],
        "expired": groups["EXPIRED"], "cancelled": groups["CANCELLED"], "other": groups["OTHER"],
        "expiring_30d": len(expiring),
        "by_contractor": contractor_rows,
        "by_type": [{"trace_type": k, "count": v} for k, v in sorted(type_counts.items(), key=lambda kv: (-kv[1], _qg_norm(kv[0])))],
        "expiring": sorted(expiring, key=lambda x: int(x.get("days_to_expire") or 0))[:20],
        "approved_by_contractor": {k: v for k, v in approved_by.items()},
    }


def _qg_trace_payload(contractor="ALL", trace_type="ALL", status="ALL", search="", plant="ALL", binder="ALL", approved_only=False, limit=3000):
    all_items = _qg_load_traces()
    filtered = _qg_trace_filter(all_items, contractor, trace_type, status, search, plant, binder, approved_only)
    return {
        "source": _qg_source_info(_qg_path(_QG_TRACE_SETTING, _QG_TRACE_DEFAULT), _QG_TRACE_DEFAULT),
        "summary": _qg_trace_summary(filtered),
        "filters": _qg_trace_filters(all_items),
        "items": filtered[:max(1, min(int(limit or 3000), 5000))],
        "total_filtered": len(filtered),
    }


_QG_NC_ALIASES = {
    "type_record": ("Tipo (RACP/RNC)", "Tipo RACP/RNC", "Tipo"),
    "emission_no": ("nº Emissão", "Nº Emissão", "Número Emissão", "Número do RNC", "Nº RNC"),
    "final_status": ("Situação Final (Status)", "Situação Final", "Status", "Situação"),
    "subject": ("Assunto",), "contractor": ("Contratada", "Empresa"),
    "description": ("Descrição Detalhada", "Descrição da NC", "Descrição"),
    "emission_date": ("Data de Emissão", "Data Emissão"),
    "response_due": ("Prazo de Resposta", "Prazo Resposta"),
    "proposed_retention": ("Valor para Retenção", "Valor a ser retido", "Valor a Reter"),
    "real_retained": ("Valor Real Retido", "Valor Retido", "Retido real"),
    "retention_status": ("Retido", "Situação Retenção", "Status Retenção"),
    "discipline": ("Disciplina",), "classification": ("Classificação",),
    "contract": ("Contrato", "Nº Contrato"), "lot": ("Lote",), "segment": ("Trecho", "Segmento", "Lote / Trecho"),
    "final_disposition": ("Disposição Final", "Decisão Final"),
    "implementation_status": ("Status Implementação", "Situação Implementação", "Implementação"),
    "effectiveness": ("Eficácia", "Status Eficácia"),
    "retained_measurement": ("Medição Retida", "Medição"), "requester": ("Solicitante",),
    "opened_by": ("Aberta por", "Aberto por"), "work_manager": ("Gestor da Obra", "Gestor obra"),
    "response_date": ("Data Resposta Contratada", "Resposta da Contratada", "Data de Resposta"),
    "supervision_validation_status": ("Validação Supervisão", "Status Validação Supervisão"),
    "ccr_validation_status": ("Validação CCR", "Status Validação CCR"),
    "close_date": ("Data Encerramento", "Encerramento"),
    "close_owner": ("Responsável Encerramento", "Responsável pelo Encerramento"),
    "action_proposal": ("Proposta de Ação da Contratada", "Proposta de Ação", "Ação Proposta"),
    "non_retention_reason": ("Justificativa da Não Retenção", "Motivo Não Retenção", "Justificativa Não Retenção"),
    "fulcrum_link": ("Link Fulcrum", "Fulcrum"),
    "anomaly_photo_link": ("Foto da Anomalia", "Link Foto Anomalia"),
    "evidence_photo_link": ("Foto da Evidência", "Link Evidência", "Evidência Fotográfica"),
}


def _qg_nc_group(status):
    s = _qg_norm(status)
    if "CANCEL" in s:
        return "CANCELLED"
    if "ENCERR" in s and "A ENCERRAR" not in s:
        return "CLOSED"
    if "TRAT" in s or "A ENCERRAR" in s or "ANDAMENTO" in s:
        return "TREATMENT"
    if "ABERT" in s:
        return "OPEN"
    return "OTHER"


def _qg_load_nc(force=True):
    p = _qg_path(_QG_NC_SETTING, _QG_NC_DEFAULT)
    if not p.exists() or not p.is_file():
        return []
    st = p.stat()
    key = str(p.resolve())
    if (
        not force
        and _qg_nc_cache["path"] == key
        and _qg_nc_cache["mtime_ns"] == st.st_mtime_ns
        and _qg_nc_cache["size"] == st.st_size
    ):
        return [dict(x) for x in _qg_nc_cache["records"]]
    read_path = _qg_stable_source_copy(p, _QG_NC_DEFAULT, "controle_nc_live")
    link_base = _qg_workbook_link_base(read_path)
    try:
        wb = load_workbook(read_path, read_only=False, data_only=True)
    except Exception as exc:
        raise ValueError("A planilha configurada de Não Conformidades não é um arquivo Excel válido.") from exc
    records = []
    try:
        actual = next((n for n in wb.sheetnames if _qg_norm(n) == "NC"), None) or next((n for n in wb.sheetnames if "NC" in _qg_norm(n)), None)
        if actual:
            ws = wb[actual]
            header_row, headers = _qg_find_header(ws, [
                ("Tipo (RACP/RNC)", "Tipo RACP/RNC", "Tipo"),
                ("nº Emissão", "Nº Emissão", "Número Emissão"),
            ], max_rows=30)
            if header_row:
                for ridx in range(header_row + 1, ws.max_row + 1):
                    vals = [ws.cell(ridx, c).value for c in range(1, ws.max_column + 1)]
                    typ = _qg_clean(_qg_pick(vals, headers, *_QG_NC_ALIASES["type_record"]))
                    emission = _qg_clean(_qg_pick(vals, headers, *_QG_NC_ALIASES["emission_no"]))
                    if _qg_norm(typ) != "RNC" or not emission:
                        continue
                    item = {"source_row": ridx, "type_record": typ, "emission_no": emission}
                    for field, aliases in _QG_NC_ALIASES.items():
                        if field in {"type_record", "emission_no", "emission_date", "response_due", "response_date", "close_date", "proposed_retention", "real_retained", "fulcrum_link", "anomaly_photo_link", "evidence_photo_link"}:
                            continue
                        item[field] = _qg_clean(_qg_pick(vals, headers, *aliases))
                    for field in ("emission_date", "response_due", "response_date", "close_date"):
                        item[field] = _qg_excel_date(_qg_pick(vals, headers, *_QG_NC_ALIASES[field]))
                    for field in ("proposed_retention", "real_retained"):
                        item[field] = _qg_money(_qg_pick(vals, headers, *_QG_NC_ALIASES[field]))
                    for field in ("fulcrum_link", "anomaly_photo_link", "evidence_photo_link"):
                        item[field] = _qg_link(
                            ws, ridx, headers, *_QG_NC_ALIASES[field], link_base=link_base,
                        ) or _qg_clean(_qg_pick(vals, headers, *_QG_NC_ALIASES[field]))
                    item["status_group"] = _qg_nc_group(item.get("final_status"))
                    item["days_overdue"] = None
                    due = str(item.get("response_due") or "")
                    if item["status_group"] in {"OPEN", "TREATMENT"} and re.fullmatch(r"\d{4}-\d{2}-\d{2}", due):
                        try:
                            item["days_overdue"] = max(0, (date.today() - date.fromisoformat(due)).days)
                        except ValueError:
                            pass
                    records.append(item)
    finally:
        wb.close()
    _qg_nc_cache.update({"path": key, "mtime_ns": st.st_mtime_ns, "size": st.st_size, "records": records})
    return [dict(x) for x in records]


def _qg_nc_filter(items, contractor="ALL", status="ALL", subject="ALL", discipline="ALL", classification="ALL", retention_status="ALL", search=""):
    cn, sn, sun, dn, cln, rn, q = map(_qg_norm, [contractor, status, subject, discipline, classification, retention_status, search])
    out = []
    for x in items:
        if cn not in {"", "ALL"} and _qg_norm(x.get("contractor")) != cn:
            continue
        if sn not in {"", "ALL"} and _qg_norm(x.get("final_status")) != sn:
            continue
        if sun not in {"", "ALL"} and _qg_norm(x.get("subject")) != sun:
            continue
        if dn not in {"", "ALL"} and _qg_norm(x.get("discipline")) != dn:
            continue
        if cln not in {"", "ALL"} and _qg_norm(x.get("classification")) != cln:
            continue
        if rn not in {"", "ALL"} and _qg_norm(x.get("retention_status")) != rn:
            continue
        if q:
            hay = _qg_norm(" ".join(str(x.get(k) or "") for k in (
                "emission_no", "contractor", "subject", "description", "segment", "discipline", "classification", "contract",
            )))
            if q not in hay:
                continue
        out.append(x)
    return out


def _qg_nc_filters(items):
    def u(k):
        return sorted({str(x.get(k) or "").strip() for x in items if str(x.get(k) or "").strip()}, key=_qg_norm)
    return {
        "contractors": u("contractor"), "statuses": u("final_status"), "subjects": u("subject"),
        "disciplines": u("discipline"), "classifications": u("classification"),
        "retention_statuses": u("retention_status"),
    }


def _qg_nc_summary(items):
    groups = Counter(x.get("status_group") or "OTHER" for x in items)
    byc = defaultdict(Counter)
    subjects = {}
    for x in items:
        c = x.get("contractor") or "Sem contratada"
        byc[c][x.get("status_group") or "OTHER"] += 1
        n = _qg_norm(x.get("subject") or "Sem assunto")
        rec = subjects.setdefault(n, {"subject": x.get("subject") or "Sem assunto", "count": 0})
        rec["count"] += 1
    contractor_rows = []
    for c in sorted(byc, key=_qg_norm):
        g = byc[c]
        total = sum(g.values())
        contractor_rows.append({
            "contractor": c, "total": total, "open": g["OPEN"], "treatment": g["TREATMENT"],
            "closed": g["CLOSED"], "cancelled": g["CANCELLED"], "other": g["OTHER"],
            "closed_pct": round(g["CLOSED"] / total * 100, 1) if total else 0,
        })
    recent_open = [x for x in items if x.get("status_group") in {"OPEN", "TREATMENT"}]
    recent_open.sort(key=lambda x: str(x.get("emission_date") or ""), reverse=True)
    return {
        "total": len(items), "open": groups["OPEN"], "treatment": groups["TREATMENT"],
        "closed": groups["CLOSED"], "cancelled": groups["CANCELLED"], "other": groups["OTHER"],
        "overdue": sum(1 for x in items if int(x.get("days_overdue") or 0) > 0),
        "proposed_retention": sum(float(x.get("proposed_retention") or 0) for x in items),
        "real_retained": sum(float(x.get("real_retained") or 0) for x in items),
        "retained_count": sum(1 for x in items if "RETID" in _qg_norm(x.get("retention_status")) and "NAO" not in _qg_norm(x.get("retention_status"))),
        "by_contractor": contractor_rows,
        "by_subject": sorted(subjects.values(), key=lambda x: (-x["count"], _qg_norm(x["subject"]))),
        "recent_open": recent_open[:20],
    }


def _qg_nc_payload(contractor="ALL", status="ALL", subject="ALL", discipline="ALL", classification="ALL", retention_status="ALL", search="", limit=3000):
    all_items = _qg_load_nc()
    filtered = _qg_nc_filter(all_items, contractor, status, subject, discipline, classification, retention_status, search)
    return {
        "source": _qg_source_info(_qg_path(_QG_NC_SETTING, _QG_NC_DEFAULT), _QG_NC_DEFAULT),
        "summary": _qg_nc_summary(filtered),
        "filters": _qg_nc_filters(all_items),
        "items": filtered[:max(1, min(int(limit or 3000), 5000))],
        "total_filtered": len(filtered),
    }


def _qg_save_upload(content: bytes, original_name: str, default: Path, setting_key: str, max_mb=60):
    if not content:
        raise ValueError("A planilha enviada está vazia.")
    if len(content) > max_mb * 1024 * 1024:
        raise ValueError(f"A planilha ultrapassa o limite de {max_mb} MB.")
    if not str(original_name or "").lower().endswith((".xlsx", ".xlsm")):
        raise ValueError("Envie uma planilha .xlsx ou .xlsm.")
    _QG_DIR.mkdir(parents=True, exist_ok=True)
    if default.exists():
        bdir = _QG_DIR / "backups"
        bdir.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        shutil.copy2(default, bdir / f"{default.stem}_{stamp}{default.suffix}")
    tmp = default.with_suffix(".tmp.xlsx")
    tmp.write_bytes(content)
    try:
        wb = load_workbook(tmp, read_only=True, data_only=True)
        wb.close()
    except Exception as exc:
        tmp.unlink(missing_ok=True)
        raise ValueError("O arquivo enviado não parece ser uma planilha Excel válida.") from exc
    os.replace(tmp, default)
    _qg_set_setting(setting_key, str(default))
    return default
