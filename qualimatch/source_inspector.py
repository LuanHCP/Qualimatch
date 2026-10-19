from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from .normalizers import norm_header


def _display(value: Any) -> Any:
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _header_candidate(values: list[Any]) -> float:
    nonempty = [value for value in values if value not in (None, "")]
    if not nonempty:
        return -1.0
    strings = [str(value).strip() for value in nonempty if isinstance(value, str)]
    normalized = [norm_header(value) for value in strings if norm_header(value)]
    unique = len(set(normalized)) / max(1, len(normalized))
    return len(strings) / len(nonempty) * 2 + unique + min(len(nonempty), 20) / 10


def _inspect_sheet(worksheet) -> dict[str, Any]:
    scan_columns = min(worksheet.max_column or 1, 200)
    scan_rows = min(worksheet.max_row or 1, 30)
    candidates: list[tuple[float, int, list[Any]]] = []
    for row_number in range(1, scan_rows + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, scan_columns + 1)]
        candidates.append((_header_candidate(values), row_number, values))

    score, header_row, raw_headers = max(candidates, key=lambda item: item[0])
    last_header = max((index for index, value in enumerate(raw_headers, 1) if value not in (None, "")), default=0)
    raw_headers = raw_headers[:last_header]
    headers = [str(value).strip() if value not in (None, "") else f"COLUNA_{index}" for index, value in enumerate(raw_headers, 1)]

    grouped: dict[str, list[str]] = {}
    for header in headers:
        grouped.setdefault(norm_header(header), []).append(header)
    duplicates = [values for key, values in grouped.items() if key and len(values) > 1]

    sample: list[list[Any]] = []
    for row_number in range(header_row + 1, min(worksheet.max_row or 0, header_row + 20) + 1):
        values = [worksheet.cell(row_number, column).value for column in range(1, len(headers) + 1)]
        if not any(value not in (None, "") for value in values):
            continue
        sample.append([_display(value) for value in values])
        if len(sample) == 5:
            break

    formula_count = 0
    for row in worksheet.iter_rows(
        min_row=1,
        max_row=min(worksheet.max_row or 1, 200),
        max_col=min(worksheet.max_column or 1, 100),
    ):
        formula_count += sum(isinstance(cell.value, str) and cell.value.startswith("=") for cell in row)

    issues: list[str] = []
    if score < 0 or not headers:
        issues.append("Não foi possível identificar um cabeçalho.")
    if duplicates:
        issues.append(f"{len(duplicates)} grupo(s) de coluna duplicada.")
    if (worksheet.max_column or 0) > 200:
        issues.append("A planilha tem mais de 200 colunas; a inspeção visual foi limitada.")
    if header_row > 10:
        issues.append(f"Cabeçalho provável encontrado apenas na linha {header_row}.")

    return {
        "name": worksheet.title,
        "state": worksheet.sheet_state,
        "rows": worksheet.max_row or 0,
        "columns": worksheet.max_column or 0,
        "header_row": header_row if headers else None,
        "headers": headers,
        "duplicate_headers": duplicates,
        "formula_count_sampled": formula_count,
        "sample": sample,
        "issues": issues,
    }


def inspect_workbook(path: str | Path) -> dict[str, Any]:
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    try:
        sheets = [_inspect_sheet(sheet) for sheet in workbook.worksheets]
    finally:
        workbook.close()

    issues = [f"{sheet['name']}: {issue}" for sheet in sheets for issue in sheet["issues"]]
    return {
        "file_name": source.name,
        "size_bytes": source.stat().st_size,
        "summary": {
            "sheets": len(sheets),
            "visible_sheets": sum(sheet["state"] == "visible" for sheet in sheets),
            "rows": sum(sheet["rows"] for sheet in sheets),
            "max_columns": max((sheet["columns"] for sheet in sheets), default=0),
            "issues": len(issues),
        },
        "issues": issues,
        "sheets": sheets,
    }
