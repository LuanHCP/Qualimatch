from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from .management import measurement_cycle_dashboard, bv_pending_workflow
from .queries import measurement_cycle
from .database import rows


def measurement_cycle_report(measurement: str, contractor: str = "ALL") -> tuple[str, bytes]:
    measurement = str(measurement or "").upper().strip()
    dash = measurement_cycle_dashboard(contractor, limit=100)
    cycle = next((x for x in dash["cycles"] if x["measurement"] == measurement), None)
    if not cycle:
        raise ValueError(f"{measurement} não encontrada na base de medições.")
    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    ws.append(["QualiMatch", f"Relatório {measurement}"])
    ws.append(["Contratada", contractor])
    ws.append(["Período", f"{cycle['measurement_start']} a {cycle['measurement_end']}"])
    ws.append([])
    ws.append(["Indicador", "Valor"])
    metrics = [
        ("Dias de execução", cycle["execution_days"]),
        ("Dias com kit", cycle["kit_days"]),
        ("Dias com certificado", cycle["certificate_days"]),
        ("Cobertura kits (%)", cycle["kit_coverage_pct"]),
        ("Cobertura certificados (%)", cycle["certificate_coverage_pct"]),
        ("Pendentes Contratada", cycle["pending_contractor"]),
        ("No prazo", cycle["on_time"]),
        ("Pendentes certificado", cycle["pending_certificate"]),
        ("Prazo médio kit (dias)", cycle["avg_kit_lead_days"]),
        ("Prazo médio certificado (dias)", cycle["avg_certificate_lead_days"]),
    ]
    for x in metrics: ws.append(list(x))
    for cell in ws[1]: cell.font=Font(bold=True,size=14)
    for cell in ws[5]: cell.font=Font(bold=True)
    ws.column_dimensions["A"].width=32; ws.column_dimensions["B"].width=28

    pend = wb.create_sheet("Pendências")
    pend.append(["Contratada","Medição","Data execução","Rodovia","Status","Kit","Data entrega","Dias"])
    for p in bv_pending_workflow(contractor, cycle["measurement_start"], cycle["measurement_end"]):
        pend.append([p.get("contractor"),p.get("measurement"),p.get("date"),p.get("road"),p.get("status"),", ".join(p.get("kit_sequences") or []),p.get("kit_delivery_date"),p.get("days_since_kit")])
    for c in pend[1]: c.font=Font(bold=True)
    for col,w in {"A":15,"B":12,"C":16,"D":18,"E":24,"F":22,"G":16,"H":12}.items(): pend.column_dimensions[col].width=w

    exec_ws = wb.create_sheet("Execuções")
    exec_ws.append(["Linha","Contratada","Data","Rodovia","Sentido","KMI","KMF","Faixa"])
    params: list[Any] = [cycle["measurement_start"],cycle["measurement_end"]]
    where=["date>=?","date<=?"]
    if contractor!="ALL": where.append("contractor=?"); params.append(contractor)
    for m in rows("SELECT line_measurement,contractor,date,road,direction,kmi,kmf,lane FROM measurements WHERE "+" AND ".join(where)+" ORDER BY date,contractor,source_row",tuple(params)):
        exec_ws.append([m.get("line_measurement"),m.get("contractor"),m.get("date"),m.get("road"),m.get("direction"),m.get("kmi"),m.get("kmf"),m.get("lane")])
    for c in exec_ws[1]: c.font=Font(bold=True)
    for row in wb.worksheets:
        for cells in row.iter_rows():
            for cell in cells: cell.alignment=Alignment(vertical="top")
    bio=BytesIO(); wb.save(bio)
    return f"QualiMatch_{measurement}_{contractor}.xlsx",bio.getvalue()
