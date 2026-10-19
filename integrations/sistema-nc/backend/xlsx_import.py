"""
Importação de planilhas Excel (.xlsx / .xlsm / .xls) para a aba NC.

Segurança: usamos openpyxl apenas em modo leitura de valores de célula
(`data_only=True`, `keep_vba=False`). O openpyxl NUNCA interpreta nem
executa código VBA/macro em nenhum modo — ele lê exclusivamente a árvore
XML dentro do pacote .xlsm/.xlsx; `keep_vba=False` apenas garante que,
mesmo ao salvar, o blob binário da macro não seria preservado. Não há
qualquer caminho de execução de macro neste fluxo.

`.xls` (formato binário antigo, pré-2007) usa um parser diferente
(`xlrd`), que não está disponível neste ambiente sandbox (mesma restrição
de rede documentada em db.py). O parser abaixo tenta `xlrd` quando
presente e, se ausente, recusa `.xls` com uma mensagem clara — a aceitação
de `.xls` funciona automaticamente assim que `xlrd` estiver instalado
(requirements.txt já o lista).
"""
import io
import re
import unicodedata
import datetime

import openpyxl

import db

try:
    import xlrd  # type: ignore
    HAVE_XLRD = True
except Exception:
    HAVE_XLRD = False


# ----------------------------------------------------------- utilidades --
def strip_accents(s):
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def normalize_header(s):
    s = strip_accents(str(s or "")).lower().strip()
    s = re.sub(r"[^a-z0-9]+", " ", s).strip()
    return s


STATUS_RULES = [
    (re.compile(r"^a encerrar"), "aguardando_verificacao"),
    (re.compile(r"cancel"), "encerrada"),
    (re.compile(r"inexistente"), "encerrada"),
    (re.compile(r"encerr"), "encerrada"),
    (re.compile(r"tratamento"), "em_analise"),
    (re.compile(r"abert"), "aguardando_resposta"),
]

# Rótulos amigáveis usados pelo próprio sistema (ex.: na exportação —
# xlsx_export.STATUS_LABEL) — comparados primeiro, com igualdade exata,
# para que um arquivo EXPORTADO pelo sistema seja REIMPORTADO com o status
# interno original preservado (o ciclo importar → exportar → reimportar
# não pode degradar/perder o status real do registro).
STATUS_LABEL_TO_KEY = {
    "rascunho": "rascunho", "emitida": "emitida", "protocolada": "protocolada",
    "aguard. resposta da contratada": "aguardando_resposta", "respondida": "respondida",
    "em análise pela concessionária": "em_analise", "disposição aprovada": "disposicao_aprovada",
    "disposição recusada": "disposicao_recusada", "em execução": "em_execucao",
    "aguard. verificação de eficácia": "aguardando_verificacao", "eficaz": "eficaz",
    "não eficaz": "nao_eficaz", "encerrada": "encerrada",
}


def normalize_status(raw):
    if not raw:
        return "aguardando_resposta"
    exact = str(raw).strip().lower()
    if exact in STATUS_LABEL_TO_KEY:
        return STATUS_LABEL_TO_KEY[exact]
    low = strip_accents(exact)
    for rx, status in STATUS_RULES:
        if rx.search(low):
            return status
    return "aguardando_resposta"


def normalize_classificacao(raw):
    if not raw:
        return None
    low = strip_accents(str(raw)).lower().strip()
    if "baix" in low or "leve" in low:
        return "leve"
    if "medi" in low or "media" in low:
        return "media"
    if "alt" in low or "grave" in low:
        return "grave"
    return low


def normalize_eficacia(raw):
    if not raw:
        return None
    low = strip_accents(str(raw)).lower().strip()
    if "nao" in low or "não" in str(raw).lower():
        return "nao_eficaz"
    if "eficaz" in low:
        return "eficaz"
    return None


def normalize_disposicao(raw):
    if not raw:
        return None
    low = strip_accents(str(raw)).lower().strip()
    if "nao aceit" in low or "recus" in low:
        return "recusada"
    if "aceit" in low or "aprov" in low:
        return "aprovada"
    return None


def coerce_date(value):
    if value is None or value == "":
        return None
    if isinstance(value, datetime.datetime):
        return value.date().isoformat()
    if isinstance(value, datetime.date):
        return value.isoformat()
    s = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def coerce_number(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).strip()
    s = re.sub(r"[^0-9,.\-]", "", s)
    if not s:
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return float(s)
    except ValueError:
        return None


def coerce_text(value):
    if value is None:
        return ""
    return str(value).strip()


# -------------------------------------------------------- header aliases -
def load_alias_map(sede_id):
    rows = db.fetchall(f"""
        SELECT campo_sistema, alias FROM header_aliases
        WHERE sede_id IS NULL OR sede_id = {db.lit(sede_id)};
    """)
    return {normalize_header(r["alias"]): r["campo_sistema"] for r in rows}


def pick_sheet(wb):
    names = wb.sheetnames
    for n in names:
        if n.strip().lower() == "nc":
            return n
    for n in names:
        nl = n.lower()
        if re.search(r"\bnc\b", nl) and not re.search(r"legenda|texto|banco", nl):
            return n
    return names[0]


# ------------------------------------------------------------ leitura ----
def read_workbook_rows(file_bytes, filename, sede_id, header_row=2):
    """Retorna (sheet_name, header_mapping, rows) onde rows é uma lista de
    dicts {'_linha': numero_da_linha_na_planilha, campo_sistema: valor_bruto}."""
    lower = filename.lower()
    if lower.endswith(".xls") and not lower.endswith(".xlsx") and not lower.endswith(".xlsm"):
        if not HAVE_XLRD:
            raise ValueError(
                "Arquivos .xls (formato binário antigo) exigem a biblioteca 'xlrd', "
                "não disponível neste ambiente. Salve a planilha como .xlsx e tente novamente."
            )
        return _read_xls(file_bytes, sede_id, header_row)

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False, read_only=True)
    sheet_name = pick_sheet(wb)
    ws = wb[sheet_name]
    alias_map = load_alias_map(sede_id)

    header_cells = next(ws.iter_rows(min_row=header_row, max_row=header_row))
    col_field = {}
    header_mapping = {}
    for idx, cell in enumerate(header_cells):
        header_text = coerce_text(cell.value)
        if not header_text:
            continue
        norm = normalize_header(header_text)
        campo = alias_map.get(norm)
        if campo:
            col_field[idx] = campo
            header_mapping[header_text] = campo

    rows = []
    for r_idx, row_cells in enumerate(ws.iter_rows(min_row=header_row + 1), start=header_row + 1):
        data = {"_linha": r_idx}
        any_value = False
        for idx, cell in enumerate(row_cells):
            if idx in col_field:
                data[col_field[idx]] = cell.value
                if cell.value not in (None, ""):
                    any_value = True
        if any_value:
            rows.append(data)
    wb.close()
    return sheet_name, header_mapping, rows


def _read_xls(file_bytes, sede_id, header_row):
    book = xlrd.open_workbook(file_contents=file_bytes)
    sheet_name = None
    for n in book.sheet_names():
        if n.strip().lower() == "nc":
            sheet_name = n
            break
    sheet_name = sheet_name or book.sheet_names()[0]
    ws = book.sheet_by_name(sheet_name)
    alias_map = load_alias_map(sede_id)
    header_vals = ws.row_values(header_row - 1)
    col_field = {}
    header_mapping = {}
    for idx, header_text in enumerate(header_vals):
        header_text = coerce_text(header_text)
        if not header_text:
            continue
        norm = normalize_header(header_text)
        campo = alias_map.get(norm)
        if campo:
            col_field[idx] = campo
            header_mapping[header_text] = campo
    rows = []
    for r_idx in range(header_row, ws.nrows):
        vals = ws.row_values(r_idx)
        data = {"_linha": r_idx + 1}
        any_value = False
        for idx, v in enumerate(vals):
            if idx in col_field:
                data[col_field[idx]] = v
                if v not in (None, ""):
                    any_value = True
        if any_value:
            rows.append(data)
    return sheet_name, header_mapping, rows


# ----------------------------------------------------- resolução aux -----
def load_contrato_aux_sheet(file_bytes, sede_sigla):
    """Lê a aba auxiliar 'Banco de Dados Contrato' para permitir vincular
    (ou pré-cadastrar) contratos referenciados na aba NC que ainda não
    existem no sistema."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, keep_vba=False, read_only=True)
    except Exception:
        return {}
    target = None
    for n in wb.sheetnames:
        if "banco" in n.lower() and "contrato" in n.lower():
            target = n
            break
    if not target:
        return {}
    ws = wb[target]
    header = [coerce_text(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = {normalize_header(h): i for i, h in enumerate(header) if h}

    def col(name):
        return idx.get(normalize_header(name))

    i_unidade, i_contrato, i_desc, i_contratada = col("Unidade"), col("Contrato"), col("Obras - Descrição Resumida"), col("Contratada")
    i_tipo, i_lote, i_gerente, i_gestor, i_ano = col("Tipo de Obra"), col("Lote"), col("Gerente"), col("Gestor da Obra"), col("Ano")
    out = {}
    for row_cells in ws.iter_rows(min_row=2):
        vals = [c.value for c in row_cells]
        if i_contrato is None or i_contrato >= len(vals):
            continue
        numero = coerce_text(vals[i_contrato])
        unidade = coerce_text(vals[i_unidade]) if i_unidade is not None else ""
        if not numero or (sede_sigla and unidade and unidade.strip().lower() != sede_sigla.strip().lower()):
            continue
        out[numero] = {
            "numero": numero,
            "descricaoObra": coerce_text(vals[i_desc]) if i_desc is not None and i_desc < len(vals) else "",
            "contratada": coerce_text(vals[i_contratada]) if i_contratada is not None and i_contratada < len(vals) else "",
            "tipoObra": coerce_text(vals[i_tipo]) if i_tipo is not None and i_tipo < len(vals) else "",
            "lote": coerce_text(vals[i_lote]) if i_lote is not None and i_lote < len(vals) else "",
            "gerente": coerce_text(vals[i_gerente]) if i_gerente is not None and i_gerente < len(vals) else "",
            "gestorObra": coerce_text(vals[i_gestor]) if i_gestor is not None and i_gestor < len(vals) else "",
            "ano": vals[i_ano] if i_ano is not None and i_ano < len(vals) else None,
        }
    wb.close()
    return out


# ------------------------------------------------------------ mapeamento -
def map_row(raw):
    """Converte uma linha bruta (campo_sistema -> valor de célula) para um
    dict normalizado e tipado, pronto para comparação/gravação."""
    numero_emissao = coerce_text(raw.get("numero_emissao"))
    unidade = coerce_text(raw.get("unidade"))
    contrato = coerce_text(raw.get("contrato"))
    return {
        "linha": raw["_linha"],
        "numeroEmissao": numero_emissao,
        "unidade": unidade,
        "contrato": contrato,
        "contratada": coerce_text(raw.get("contratada")),
        "dataEmissao": coerce_date(raw.get("data_emissao")),
        "situacaoRaw": coerce_text(raw.get("situacao_final")) or coerce_text(raw.get("status")),
        "tipoRaw": coerce_text(raw.get("tipo_core")),
        "categoria": coerce_text(raw.get("assunto")) or coerce_text(raw.get("categoria")),
        "item": coerce_text(raw.get("item")),
        "subitem": coerce_text(raw.get("subitem")),
        "descricao": coerce_text(raw.get("descricao")),
        "disciplina": coerce_text(raw.get("disciplina")),
        "classificacao": normalize_classificacao(raw.get("classificacao")),
        "valorRetencaoProposto": coerce_number(raw.get("valor_retencao_proposto")) or coerce_number(raw.get("valor")),
        "valorRetencaoEfetivo": coerce_number(raw.get("valor_retencao_efetivo")),
        "prazoResposta": coerce_date(raw.get("prazo_resposta_dias")),
        "propostaAcao": coerce_text(raw.get("proposta_acao")),
        "parecerSupervisao": coerce_text(raw.get("status_validacao_supervisao")),
        "eficaciaRaw": coerce_text(raw.get("eficacia_solucao")),
        "dataEncerramento": coerce_date(raw.get("data_encerramento")),
        "rodovia": coerce_text(raw.get("rodovia")),
        "kmInicial": coerce_number(raw.get("km_inicial")),
        "kmFinal": coerce_number(raw.get("km_final")),
    }
