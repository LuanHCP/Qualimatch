from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

import qualimatch.management as mg
import qualimatch.reporting as rp
from qualimatch.auth import (
    can_edit,
    create_or_update_admin,
    create_registration,
    get_user,
    set_user_permissions,
    set_user_status,
)
from qualimatch.database import init_db, transaction, row


def test_viewer_editor_admin_permission_levels(tmp_path: Path):
    db = tmp_path / 'v380.db'
    init_db(db)
    admin = create_or_update_admin('Luan', 'admin@example.com', 'Senha123', db)
    user = create_registration('Usuário', 'user@example.com', 'Senha456', db)
    user = set_user_status(user['id'], 'APPROVE', admin['id'], db)
    user = set_user_permissions(
        user['id'], True, True, False, admin['id'], db,
        access_arena=True, access_pavidez=False, permission_level='VIEWER'
    )
    assert user['permission_level'] == 'VIEWER'
    assert not can_edit(user)
    user = set_user_permissions(
        user['id'], True, True, False, admin['id'], db,
        access_arena=True, access_pavidez=False, permission_level='EDITOR'
    )
    assert user['permission_level'] == 'EDITOR'
    assert can_edit(user)
    assert can_edit(get_user(admin['id'], db))


def test_bv_pending_workflow_classifies_contractor_and_certificate(monkeypatch):
    monkeypatch.setattr(mg, 'pending_days', lambda contractor, start, end: [
        {'contractor':'ARENA','date':'2026-08-10','measurement':'M13','road':'BR-376'},
        {'contractor':'ARENA','date':'2026-08-11','measurement':'M14','road':'BR-376'},
        {'contractor':'ARENA','date':'2026-08-12','measurement':'M14','road':'BR-376'},
    ])
    def fake_rows(sql, params=(), path=None):
        if 'FROM kit_rows' in sql:
            return [
                {'contractor':'ARENA','application_date':'2026-08-11','receipt_date':'2026-08-15','status':'OK','seq':'101'},
                {'contractor':'ARENA','application_date':'2026-08-12','receipt_date':'2026-08-19','status':'OK','seq':'102'},
            ]
        return []
    monkeypatch.setattr(mg, 'rows', fake_rows)
    monkeypatch.setattr(mg, '_today', lambda: mg.date(2026,8,19))
    out = mg.bv_pending_workflow('ARENA')
    assert [x['status'] for x in out] == ['PENDENTE CONTRATADA','PENDENTE CERTIFICADO','NO PRAZO']
    assert out[2]['kit_sequences'] == ['102']


def test_measurement_cycle_dashboard_summarizes_m14(monkeypatch):
    def fake_rows(sql, params=(), path=None):
        if 'SELECT DISTINCT contractor,date FROM measurements' in sql:
            return [
                {'contractor':'ARENA','date':'2026-08-11'},
                {'contractor':'ARENA','date':'2026-08-12'},
            ]
        if 'SELECT contractor,application_date,receipt_date FROM kit_rows' in sql:
            return [
                {'contractor':'ARENA','application_date':'2026-08-11','receipt_date':'2026-08-12'},
            ]
        if 'SELECT id,contractor,service_date,extraction_date,raw_json FROM certificates' in sql:
            return [
                {'id':1,'contractor':'ARENA','service_date':'2026-08-11','extraction_date':'2026-08-14','raw_json':'{}'},
            ]
        if 'certificate_samples' in sql:
            return []
        return []
    monkeypatch.setattr(mg, 'rows', fake_rows)
    monkeypatch.setattr(mg, 'bv_pending_workflow', lambda contractor: [
        {'date':'2026-08-12','status':'PENDENTE CONTRATADA'}
    ])
    out = mg.measurement_cycle_dashboard('ARENA')
    m14 = next(x for x in out['cycles'] if x['measurement']=='M14')
    assert m14['execution_days'] == 2
    assert m14['kit_days'] == 1
    assert m14['certificate_days'] == 1
    assert m14['pending_contractor'] == 1
    assert m14['kit_coverage_pct'] == 50.0
    assert m14['certificate_coverage_pct'] == 50.0
    assert m14['avg_kit_lead_days'] == 1.0
    assert m14['avg_certificate_lead_days'] == 3.0


def test_traceability_returns_measurement_kit_and_certificate(monkeypatch):
    monkeypatch.setattr(mg, 'row', lambda sql, params=(), path=None: {
        'id':7,'contractor':'PAVIDEZ','date':'2026-08-15','road':'BR-376','line_measurement':'3646'
    } if 'FROM measurements' in sql else None)
    calls = {'n':0}
    def fake_rows(sql, params=(), path=None):
        if 'FROM kit_rows' in sql:
            return [{'id':3,'seq':'55','contractor':'PAVIDEZ','application_date':'2026-08-15','status':'OK'}]
        if 'FROM matches' in sql:
            return [{'id':9,'number':'BI330-999','contractor':'PAVIDEZ','routine_result':'APROVADO'}]
        return []
    monkeypatch.setattr(mg, 'rows', fake_rows)
    out = mg.traceability(7)
    assert out['measurement']['measurement'] == 'M14'
    assert out['kits'][0]['seq'] == '55'
    assert out['certificates'][0]['number'] == 'BI330-999'


def test_global_search_respects_contractor_scope(monkeypatch):
    seen = []
    def fake_rows(sql, params=(), path=None):
        seen.append((sql, params))
        if 'FROM kit_rows' in sql:
            assert params[0] == 'ARENA'
            return [{'id':1,'contractor':'ARENA','seq':'777','application_date':'2026-08-15','road':'BR-376','status':'OK'}]
        if 'FROM kit_documents' in sql:
            assert params[0] == 'ARENA'
            return []
        return []
    monkeypatch.setattr(mg, 'rows', fake_rows)
    user = {'role':'USER','status':'APPROVED','access_bv':0,'access_arena':1,'access_pavidez':0}
    out = mg.global_search('777', user)
    assert len(out) == 1
    assert out[0]['contractor'] == 'ARENA'
    assert all('PAVIDEZ' not in str(params) for _, params in seen)


def test_change_history_undo_kit_on_temp_database(tmp_path: Path):
    db = tmp_path / 'undo.db'
    init_db(db)
    admin = create_or_update_admin('Luan', 'admin@example.com', 'Senha123', db)
    with transaction(db) as con:
        con.execute("INSERT INTO kit_rows(contractor,source_row,seq,status) VALUES(?,?,?,?)", ('ARENA',1,'10','PENDENTE'))
    kit = row('SELECT * FROM kit_rows WHERE contractor=? AND source_row=?', ('ARENA',1), db)
    before = dict(kit)
    with transaction(db) as con:
        con.execute("UPDATE kit_rows SET status='OK' WHERE id=?", (kit['id'],))
    after = row('SELECT * FROM kit_rows WHERE id=?', (kit['id'],), db)
    cid = mg.record_change(admin['id'], 'KIT_EDIT', 'kit', kit['id'], before, after, db)
    assert cid > 0
    result = mg.undo_change(cid, admin['id'], db)
    restored = row('SELECT status FROM kit_rows WHERE id=?', (kit['id'],), db)
    assert result['ok'] is True
    assert restored['status'] == 'PENDENTE'
    changes = mg.list_changes(10, db)
    assert changes[0]['undone_at']


def test_measurement_cycle_report_generates_xlsx(monkeypatch):
    monkeypatch.setattr(rp, 'measurement_cycle_dashboard', lambda contractor, limit=100: {'cycles':[{
        'measurement':'M14','measurement_start':'2026-08-11','measurement_end':'2026-09-10',
        'execution_days':2,'kit_days':1,'certificate_days':1,'kit_coverage_pct':50.0,
        'certificate_coverage_pct':50.0,'pending_contractor':1,'on_time':0,'pending_certificate':0,
        'avg_kit_lead_days':1.0,'avg_certificate_lead_days':3.0,
    }]})
    monkeypatch.setattr(rp, 'bv_pending_workflow', lambda contractor, start, end: [{
        'contractor':'ARENA','measurement':'M14','date':'2026-08-12','road':'BR-376',
        'status':'PENDENTE CONTRATADA','kit_sequences':[],'kit_delivery_date':None,'days_since_kit':None,
    }])
    monkeypatch.setattr(rp, 'rows', lambda sql, params=(): [{
        'line_measurement':'3646','contractor':'ARENA','date':'2026-08-12','road':'BR-376',
        'direction':'SUL','kmi':300000,'kmf':301000,'lane':'1'
    }])
    name, content = rp.measurement_cycle_report('M14','ARENA')
    wb = load_workbook(BytesIO(content), data_only=False)
    assert name == 'QualiMatch_M14_ARENA.xlsx'
    assert wb.sheetnames == ['Resumo','Pendências','Execuções']
    assert wb['Resumo']['B1'].value == 'Relatório M14'
    assert wb['Pendências'].max_row == 2


def test_v380_ui_exposes_management_features():
    root = Path(__file__).resolve().parents[1]
    html = (root/'frontend'/'index.html').read_text(encoding='utf-8')
    js = (root/'frontend'/'static'/'app.js').read_text(encoding='utf-8')
    assert 'Gestão Mxx' in html
    assert 'Revisões' in html
    assert 'Saúde do sistema' in html
    assert 'global-search-input' in html
    assert 'notification-btn' in html
    assert 'Ver fluxo' in js
    assert '/api/reports/measurement/' in js
    assert 'permission_level' in js
    assert 'Desenvolvido por Luan' in html


def test_sqlite_backup_restore_roundtrip(tmp_path: Path, monkeypatch):
    import qualimatch.database as dbm
    db = tmp_path / 'main.db'
    backups = tmp_path / 'backups'
    init_db(db)
    with transaction(db) as con:
        con.execute("INSERT INTO settings(key,value) VALUES(?,?)", ('restore_test','original'))
    # Faz com que as rotinas sem path trabalhem sobre o DB temporário deste teste.
    monkeypatch.setattr(dbm, 'DB_PATH', db)
    monkeypatch.setattr(dbm, 'BACKUP_DIR', backups)
    monkeypatch.setattr(dbm, 'USE_POSTGRES', False)
    backup_name = dbm.automatic_backup('teste', 0)
    with transaction(db) as con:
        con.execute("UPDATE settings SET value='alterado' WHERE key='restore_test'")
    assert row("SELECT value FROM settings WHERE key='restore_test'", path=db)['value'] == 'alterado'
    result = dbm.restore_backup(backup_name)
    assert result['ok'] is True
    assert row("SELECT value FROM settings WHERE key='restore_test'", path=db)['value'] == 'original'
    assert (backups / result['safety_backup']).exists()
